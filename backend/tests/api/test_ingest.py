"""Unit tests for Exam Ingestion API Endpoints (TASK-202)."""

import io
from unittest.mock import MagicMock, patch

import docx
import httpcore
import httpx
import pytest
from postgrest.exceptions import APIError

MOCK_SESSION_ID = "11111111-1111-1111-1111-111111111111"


@pytest.fixture
def auth_headers() -> dict[str, str]:
    return {"Authorization": "Bearer mock-valid-jwt-token"}


@pytest.fixture
def mock_auth_user(mock_user_response):
    mock_supabase = MagicMock()
    mock_supabase.auth.get_user.return_value = mock_user_response
    return mock_supabase


def create_sample_docx_bytes() -> bytes:
    """Helper to generate a valid .docx exam file."""
    doc = docx.Document()
    doc.add_paragraph("Câu 1: Giải phương trình bậc hai $x^2 - 1 = 0$.")
    doc.add_paragraph("A. $x = \\pm 1$  B. $x = 1$  C. $x = -1$  D. $x = 0$")
    doc.add_paragraph("Câu 2: Đạo hàm của $y = 2x$ là:")
    doc.add_paragraph("A. $2$  B. $1$  C. $x$  D. $0$")
    doc.add_paragraph("BẢNG ĐÁP ÁN:")
    doc.add_paragraph("1.A 2.A")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_docx_with_table_and_omml() -> bytes:
    """Generate docx with 1 question outside table and 1 question with OMML in table (Suggestion 1)."""
    from docx.oxml import parse_xml

    doc = docx.Document()
    doc.add_paragraph("Câu 1: Giải phương trình bậc hai $x^2 - 1 = 0$.")
    doc.add_paragraph("A. $x = 1$  B. $x = 2$  C. $x = 3$  D. $x = 4$")

    t = doc.add_table(rows=2, cols=1)
    cell_stem = t.cell(0, 0)
    cell_stem.paragraphs[0].text = "Câu 2: Đạo hàm của hàm số sau:"
    omml_xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:oMath><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e>"
        "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath></w:p>"
    )
    p_omml = cell_stem.add_paragraph("")
    p_omml._p.append(parse_xml(omml_xml))

    cell_opts = t.cell(1, 0)
    cell_opts.paragraphs[0].text = "A. 2  B. 1  C. 0  D. -1"

    doc.add_paragraph("BẢNG ĐÁP ÁN:")
    doc.add_paragraph("1.A  2.B")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_docx_with_empty_option() -> bytes:
    """Generate docx with option A having empty content (Suggestion 2)."""
    doc = docx.Document()
    doc.add_paragraph("Câu 1: Chọn giá trị đúng:")
    doc.add_paragraph("A.")
    doc.add_paragraph("B. 2  C. 3  D. 4")
    doc.add_paragraph("BẢNG ĐÁP ÁN:")
    doc.add_paragraph("1.B")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 1. Unauthenticated & Authorization Tests
# ---------------------------------------------------------------------------


def test_ingest_file_unauthenticated(client):
    """POST /api/sessions/{id}/ingest-file without token returns 401."""
    response = client.post(
        f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
        files={
            "file": (
                "test.docx",
                b"dummy content",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 401


def test_ingest_text_unauthenticated(client):
    """POST /api/sessions/{id}/ingest-text without token returns 401."""
    response = client.post(
        f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
        json={"raw_text": "Câu 1: Test"},
    )
    assert response.status_code == 401


# ---------------------------------------------------------------------------
# 2. State Precondition Tests (Status != draft returns 409)
# ---------------------------------------------------------------------------


def test_ingest_text_session_not_in_draft_returns_409(client, mock_auth_user, auth_headers):
    """Ingestion rejected with 409 Conflict if session status is already 'blueprinted' or 'generating'."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "generating"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": "Câu 1: Nội dung câu 1\nA. 1 B. 2 C. 3 D. 4\nĐÁP ÁN: 1.A"},
            headers=auth_headers,
        )
    assert response.status_code == 409
    assert "draft" in response.json()["detail"]


# ---------------------------------------------------------------------------
# 3. File Security Tests (SEC-04)
# ---------------------------------------------------------------------------


def test_ingest_file_unsupported_pdf_returns_400(client, mock_auth_user, auth_headers):
    """Reject PDF upload with HTTP 400 Bad Request."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={"file": ("exam.pdf", b"%PDF-1.4...", "application/pdf")},
            headers=auth_headers,
        )
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]


# ---------------------------------------------------------------------------
# 4. Zero Mutation & In-Memory Pre-validation (HTTP 422)
# ---------------------------------------------------------------------------


def test_ingest_text_missing_answer_returns_422(client, mock_auth_user, auth_headers):
    """Missing answer key triggers HTTP 422 and database is untouched."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    raw_text = """
### Câu 1: Cho hàm số $f(x) = x^2$.
A. 1  B. 2  C. 3  D. 4

### Câu 2: Không có đáp án đúng
A. 1  B. 2  C. 3  D. 4

BẢNG ĐÁP ÁN:
1.A
"""
    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 422
    data = response.json()["detail"]
    assert any(
        err["question_no"] == 2 and err["error_type"] == "MISSING_ANSWER" for err in data["errors"]
    )

    # Ensure NO insert or delete called on database (Database untouched!)
    mock_db.table.return_value.insert.assert_not_called()
    mock_db.table.return_value.delete.assert_not_called()


def test_ingest_file_docx_omml_in_table_returns_422_no_rpc_no_storage(
    client, mock_auth_user, auth_headers
):
    """Suggestion 1: DOCX file containing OMML in a table must return 422 and call neither RPC nor Storage."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    docx_bytes = create_docx_with_table_and_omml()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "test_omml_table.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 422
    data = response.json()["detail"]
    assert any(err["error_type"] == "UNPARSED_OMML" for err in data["errors"])

    # Guaranteed Zero Mutation & no Storage upload
    mock_db.rpc.assert_not_called()
    mock_db.storage.from_.assert_not_called()
    mock_db.table.return_value.insert.assert_not_called()
    mock_db.table.return_value.delete.assert_not_called()


def test_ingest_text_empty_option_content_returns_422_no_rpc_no_storage(
    client, mock_auth_user, auth_headers
):
    """Suggestion 2: Text ingestion with empty option content must return 422 and call neither RPC nor Storage."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    raw_text = (
        "### Câu 1: Tìm x biết x > 0:\n"
        "A. \n"
        "B. 2\n"
        "C. 3\n"
        "D. 4\n\n"
        "BẢNG ĐÁP ÁN:\n"
        "1.B\n"
    )
    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 422
    data = response.json()["detail"]
    assert any(err["error_type"] == "INVALID_FORMAT" for err in data["errors"])

    mock_db.rpc.assert_not_called()
    mock_db.storage.from_.assert_not_called()
    mock_db.table.return_value.insert.assert_not_called()
    mock_db.table.return_value.delete.assert_not_called()


def test_ingest_file_empty_option_content_returns_422_no_rpc_no_storage(
    client, mock_auth_user, auth_headers
):
    """Suggestion 2: DOCX ingestion with empty option content must return 422 and call neither RPC nor Storage."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    docx_bytes = create_docx_with_empty_option()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "empty_opt.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 422
    data = response.json()["detail"]
    assert any(err["error_type"] == "INVALID_FORMAT" for err in data["errors"])

    mock_db.rpc.assert_not_called()
    mock_db.storage.from_.assert_not_called()
    mock_db.table.return_value.insert.assert_not_called()
    mock_db.table.return_value.delete.assert_not_called()


def test_ingest_text_question_stems_not_misidentified_as_answer_key(
    client, mock_auth_user, auth_headers
):
    """Suggestion 3: Text starting with 'Câu 1: A, B là...' must not be cut off and should return 200."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 3}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "A"},
        {"id": "q-2", "question_no": 2, "correct_key": "B"},
        {"id": "q-3", "question_no": 3, "correct_key": "C"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    raw_text = """Câu 1: A, B là hai biến cố độc lập trong không gian mẫu.
*A.* 0.5  B. 0.2  C. 0.3  D. 0.1

Câu 2: A, B là hai ma trận vuông cấp hai khả nghịch.
A. 1  *B.* 2  C. 3  D. 4

Câu 3: A, B là hai tập hợp các điểm trên mặt phẳng toạ độ.
A. 1  B. 2  *C.* 3  D. 4
"""
    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 3


def test_ingest_text_latex_enumerate_options_parsed_correctly(client, mock_auth_user, auth_headers):
    """Suggestion 4: LaTeX with question headers outside enumerate options returns 200 and parses 2 questions."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "A"},
        {"id": "q-2", "question_no": 2, "correct_key": "A"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    raw_latex = r"""\textbf{Câu 1:} Cho hàm số $y=f(x)$ liên tục trên $\mathbb{R}$.
\begin{enumerate}
    \item $x = 1$
    \item $x = 2$
    \item $x = 3$
    \item $x = 4$
\end{enumerate}

\textbf{Câu 2:} Tính đạo hàm của hàm số $y=2x$.
\begin{enumerate}[(A)]
    \item $y' = 2$
    \item $y' = 1$
    \item $y' = 0$
    \item $y' = -2$
\end{enumerate}

BẢNG ĐÁP ÁN:
1.A  2.A
"""
    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_latex},
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 2


# ---------------------------------------------------------------------------
# 5. Success Ingestion (HTTP 200) & Atomic Replacement
# ---------------------------------------------------------------------------


def test_ingest_text_success_200(client, mock_auth_user, auth_headers):
    """Valid text ingestion replaces questions atomically and returns 200."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    # Mock RPC replacement success
    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 1}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    # Mock query returning inserted questions
    mock_q_res = MagicMock()
    mock_q_res.data = [
        {
            "id": "q-1",
            "session_id": MOCK_SESSION_ID,
            "question_no": 1,
            "raw_content": "Giải phương trình $x + 1 = 2$",
            "options": {"A": "$x = 1$", "B": "$x = 2$", "C": "$x = 3$", "D": "$x = 4$"},
            "correct_key": "A",
            "has_missing_answer": False,
            "has_security_warning": False,
        }
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    raw_text = """
### Câu 1: Giải phương trình $x + 1 = 2$
A. $x = 1$
B. $x = 2$
C. $x = 3$
D. $x = 4$

BẢNG ĐÁP ÁN:
1.A
"""
    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 1
    assert res_data["session_id"] == MOCK_SESSION_ID
    assert len(res_data["questions"]) == 1


def test_ingest_file_docx_success_200(client, mock_auth_user, auth_headers):
    """Upload valid DOCX file, saves to storage, and replaces questions atomically."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    # Mock storage upload
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Mock RPC
    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    # Mock query returning inserted questions
    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "A"},
        {"id": "q-2", "question_no": 2, "correct_key": "A"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 2
    assert "successfully" in res_data["message"]


def test_ingest_file_docx_multi_column_questions_success_200(client, mock_auth_user, auth_headers):
    """Review Round 2 [P2]: Multi-column question table DOCX uploaded via API succeeds (200) with 2 questions."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "A"},
        {"id": "q-2", "question_no": 2, "correct_key": "B"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    # Build docx with 1-row, 2-cell table containing Câu 1 and Câu 2
    doc = docx.Document()
    tbl = doc.add_table(rows=1, cols=2)
    c0 = tbl.rows[0].cells[0]
    c0.paragraphs[0].text = "Câu 1: Tính 1+1?"
    c0.add_paragraph("A. 2\nB. 3\nC. 4\nD. 5")

    c1 = tbl.rows[0].cells[1]
    c1.paragraphs[0].text = "Câu 2: Tính 2+2?"
    c1.add_paragraph("A. 3\nB. 4\nC. 5\nD. 6")

    # Separated answer table (1 | A | 2 | B)
    tbl_ans = doc.add_table(rows=1, cols=4)
    tbl_ans.rows[0].cells[0].text = "1"
    tbl_ans.rows[0].cells[1].text = "A"
    tbl_ans.rows[0].cells[2].text = "2"
    tbl_ans.rows[0].cells[3].text = "B"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "multicolumn.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 2

    # Verify RPC payload preserved individual questions and correct option assignments
    rpc_call_args = mock_db.rpc.call_args[0][1]["p_questions"]
    assert len(rpc_call_args) == 2
    assert rpc_call_args[0]["question_no"] == 1
    assert "Tính 1+1?" in rpc_call_args[0]["raw_content"]
    assert rpc_call_args[0]["options"]["B"] == "3"
    assert rpc_call_args[0]["correct_key"] == "A"

    assert rpc_call_args[1]["question_no"] == 2
    assert "Tính 2+2?" in rpc_call_args[1]["raw_content"]
    assert rpc_call_args[1]["options"]["B"] == "4"
    assert rpc_call_args[1]["correct_key"] == "B"


def test_ingest_file_docx_grid_answer_table_success_200(client, mock_auth_user, auth_headers):
    """Review Round 2 [P2]: DOCX with 2-row grid answer table succeeds (200) without 422 MISSING_ANSWER."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "C"},
        {"id": "q-2", "question_no": 2, "correct_key": "D"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    doc = docx.Document()
    doc.add_paragraph("Câu 1: Câu hỏi toán số một.")
    doc.add_paragraph("A. 1\nB. 2\nC. 3\nD. 4")
    doc.add_paragraph("Câu 2: Câu hỏi toán số hai.")
    doc.add_paragraph("A. 1\nB. 2\nC. 3\nD. 4")

    # 2-row answer grid table
    tbl_ans = doc.add_table(rows=2, cols=3)
    tbl_ans.rows[0].cells[0].text = "Câu"
    tbl_ans.rows[0].cells[1].text = "1"
    tbl_ans.rows[0].cells[2].text = "2"
    tbl_ans.rows[1].cells[0].text = "Đ/A"
    tbl_ans.rows[1].cells[1].text = "C"
    tbl_ans.rows[1].cells[2].text = "D"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "exam_grid_ans.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 2


def test_ingest_file_rpc_failure_returns_500_zero_mutation(client, mock_auth_user, auth_headers):
    """If RPC fails with confirmed DB error (APIError), return 500 confirming rollback without client mutations."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    # Mock storage upload
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Mock RPC failure with confirmed DB rollback (APIError from PostgreSQL/PostgREST)
    mock_db.rpc.return_value.execute.side_effect = APIError(
        {"code": "P0001", "message": "PostgreSQL transaction error: constraint violation"}
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    assert "Database atomic replacement failed and was rolled back" in response.json()["detail"]
    # Ensure NO client-side delete was called on question_blueprints or source_questions
    assert not mock_db.table.return_value.delete.called


def test_ingest_file_rpc_read_timeout_empty_message_returns_504(
    client, mock_auth_user, auth_headers
):
    """Real httpx.ReadTimeout with empty string message must return 504 per docs, without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Real httpx.ReadTimeout exception with empty string message
    mock_db.rpc.return_value.execute.side_effect = httpx.ReadTimeout("")

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 504
    detail = response.json()["detail"]
    assert "uncertain due to network timeout" in detail
    assert "was rolled back" not in detail
    assert "Session state was not modified" not in detail
    assert "Please verify session state" in detail


def test_ingest_file_rpc_read_error_empty_message_returns_504(client, mock_auth_user, auth_headers):
    """Real httpx.ReadError with empty message must return 504 warning uncertain status without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Real httpx.ReadError exception with empty string message
    mock_db.rpc.return_value.execute.side_effect = httpx.ReadError("")

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 504
    detail = response.json()["detail"]
    assert "uncertain due to network communication error" in detail
    assert "was rolled back" not in detail
    assert "Please verify session state" in detail


def test_ingest_file_rpc_remote_protocol_error_lost_response_returns_504(
    client, mock_auth_user, auth_headers
):
    """RemoteProtocolError ('Server disconnected without sending a response') must return 504 without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Real httpx.RemoteProtocolError for dropped response after potential commit
    mock_db.rpc.return_value.execute.side_effect = httpx.RemoteProtocolError(
        "Server disconnected without sending a response."
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 504
    detail = response.json()["detail"]
    assert "uncertain due to network communication error" in detail
    assert "was rolled back" not in detail
    assert "Please verify session state" in detail


def test_ingest_file_rpc_httpcore_remote_protocol_error_returns_504(
    client, mock_auth_user, auth_headers
):
    """httpcore.RemoteProtocolError directly raised must also be recognized as network communication failure."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Real httpcore.RemoteProtocolError
    mock_db.rpc.return_value.execute.side_effect = httpcore.RemoteProtocolError(
        "Server disconnected without sending a response."
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 504
    detail = response.json()["detail"]
    assert "uncertain due to network communication error" in detail
    assert "was rolled back" not in detail
    assert "Please verify session state" in detail


def test_ingest_file_rpc_unexpected_failure_without_db_evidence_returns_500_no_rollback_assertion(
    client, mock_auth_user, auth_headers
):
    """Unverified exceptions without DB evidence return 500 without falsely claiming transaction was rolled back."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Unexpected Python runtime fault without DB confirmation
    mock_db.rpc.return_value.execute.side_effect = RuntimeError(
        "Worker process memory allocation fault"
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    detail = response.json()["detail"]
    assert "was rolled back" not in detail
    assert "Please verify session state before retrying." in detail


def test_ingest_text_rpc_read_timeout_empty_message_returns_504(
    client, mock_auth_user, auth_headers
):
    """ingest-text endpoint also returns 504 uncertain status when RPC raises httpx.ReadTimeout("")."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.rpc.return_value.execute.side_effect = httpx.ReadTimeout("")

    raw_text = "Câu 1: Cho $x = 1$. Giá trị của $x$ là:\nA. 1  B. 2  C. 3  D. 4\nBẢNG ĐÁP ÁN: 1.A"

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 504
    detail = response.json()["detail"]
    assert "uncertain due to network timeout" in detail
    assert "was rolled back" not in detail
    assert "Please verify session state" in detail


def test_ingest_file_rpc_gateway_502_html_returns_500_neutral_message(
    client, mock_auth_user, auth_headers
):
    """Gateway 502 Bad Gateway HTML wrapped in APIError must return 500 neutral message without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Exact format generated by postgrest-py generate_default_error_message(r) on gateway 502 HTML
    mock_db.rpc.return_value.execute.side_effect = APIError(
        {
            "code": 502,
            "message": "JSON could not be generated",
            "hint": "Refer to full message for details",
            "details": "<html><head><title>502 Bad Gateway</title></head><body><center><h1>502 Bad Gateway</h1></center></body></html>",
        }
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    detail = response.json()["detail"]
    assert "was rolled back" not in detail
    assert "Session state was not modified" not in detail
    assert "Please verify session state before retrying." in detail
    assert not mock_db.table.return_value.delete.called


def test_ingest_file_rpc_gateway_504_html_returns_uncertain_neutral_message(
    client, mock_auth_user, auth_headers
):
    """Gateway 504 Gateway Timeout HTML wrapped in APIError must return uncertain status without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # Exact format generated by postgrest-py on gateway 504 HTML
    mock_db.rpc.return_value.execute.side_effect = APIError(
        {
            "code": 504,
            "message": "JSON could not be generated",
            "hint": "Refer to full message for details",
            "details": "<html><head><title>504 Gateway Timeout</title></head><body><center><h1>504 Gateway Timeout</h1></center></body></html>",
        }
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    # May return 504 (recognized timeout) or 500 neutral; neither must affirm rollback
    assert response.status_code in (500, 504)
    detail = response.json()["detail"]
    assert "was rolled back" not in detail
    assert "Session state was not modified" not in detail
    assert "Please verify session state" in detail
    assert not mock_db.table.return_value.delete.called


def test_ingest_file_rpc_non_json_unknown_code_returns_500_neutral_message(
    client, mock_auth_user, auth_headers
):
    """APIError with unknown/non-PostgreSQL code and non-JSON body returns 500 neutral message without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    mock_db.rpc.return_value.execute.side_effect = APIError(
        {
            "code": "PROXY_CONN_ERR",
            "message": "JSON could not be generated",
            "hint": "Refer to full message for details",
            "details": "Connection to upstream server terminated unexpectedly",
        }
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    detail = response.json()["detail"]
    assert "was rolled back" not in detail
    assert "Session state was not modified" not in detail
    assert "Please verify session state before retrying." in detail
    assert not mock_db.table.return_value.delete.called


def test_ingest_file_rpc_apierror_none_code_returns_500_neutral_message(
    client, mock_auth_user, auth_headers
):
    """APIError with None code returns 500 neutral message without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    mock_db.rpc.return_value.execute.side_effect = APIError(
        {
            "code": None,
            "message": "Unspecified database proxy error",
            "hint": None,
            "details": None,
        }
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    detail = response.json()["detail"]
    assert "was rolled back" not in detail
    assert "Please verify session state before retrying." in detail
    assert not mock_db.table.return_value.delete.called


def test_ingest_file_rpc_sqlstate_08007_transaction_resolution_unknown_returns_500_neutral_message(
    client, mock_auth_user, auth_headers
):
    """PostgreSQL SQLSTATE 08007 (transaction resolution unknown) must return 500 neutral message without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # SQLSTATE 08007: Transaction resolution unknown (PostgreSQL Class 08 Connection Exception)
    mock_db.rpc.return_value.execute.side_effect = APIError(
        {
            "code": "08007",
            "message": "transaction resolution unknown",
            "hint": None,
            "details": "Connection dropped after query submission; commit status uncertain",
        }
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    detail = response.json()["detail"]
    assert "was rolled back" not in detail
    assert "Session state was not modified" not in detail
    assert "Please verify session state before retrying." in detail
    assert not mock_db.table.return_value.delete.called


def test_ingest_file_rpc_sqlstate_08006_connection_failure_returns_500_neutral_message(
    client, mock_auth_user, auth_headers
):
    """PostgreSQL SQLSTATE 08006 (connection failure) must return 500 neutral message without claiming rollback."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # SQLSTATE 08006: Connection failure (PostgreSQL Class 08 Connection Exception)
    mock_db.rpc.return_value.execute.side_effect = APIError(
        {
            "code": "08006",
            "message": "connection failure",
            "hint": None,
            "details": "Connection to PostgreSQL server lost",
        }
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    detail = response.json()["detail"]
    assert "was rolled back" not in detail
    assert "Session state was not modified" not in detail
    assert "Please verify session state before retrying." in detail
    assert not mock_db.table.return_value.delete.called


def test_ingest_file_rpc_sqlstate_23505_unique_violation_returns_500_confirmed_rollback(
    client, mock_auth_user, auth_headers
):
    """PostgreSQL SQLSTATE 23505 (unique_violation) confirms rollback and returns 500."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]
    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    # SQLSTATE 23505: Unique constraint violation (confirmed engine rollback)
    mock_db.rpc.return_value.execute.side_effect = APIError(
        {
            "code": "23505",
            "message": "duplicate key value violates unique constraint",
            "hint": None,
            "details": "Key (session_id, question_no)=(...) already exists.",
        }
    )

    docx_bytes = create_sample_docx_bytes()

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "toan12.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 500
    detail = response.json()["detail"]
    assert "Database atomic replacement failed and was rolled back." in detail
    assert not mock_db.table.return_value.delete.called


def test_ingest_text_preserves_probability_formulas_full_options(
    client, mock_auth_user, auth_headers
):
    """Regression test: ingest-text preserves probability formulas P(A)=0.1 without truncating to '= 0.1$'."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 1}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {
            "id": "q-1",
            "question_no": 1,
            "raw_content": "Câu 1: Cho các biến cố A, B.",
            "options": {"A": "P(A)=0.1", "B": "P(B)=0.2", "C": "P(C)=0.3", "D": "P(D)=0.4"},
            "correct_option": "A",
            "correct_key": "A",
        }
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    raw_text = """### Câu 1: Cho các biến cố A, B. Khẳng định nào sau đây đúng?
A. P(A)=0.1
B. P(B)=0.2
C. P(C)=0.3
D. P(D)=0.4

BẢNG ĐÁP ÁN:
1.A"""

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 200
    data = response.json()
    assert data["total_questions"] == 1
    q = data["questions"][0]
    assert q["options"]["A"] == "P(A)=0.1"
    assert q["options"]["B"] == "P(B)=0.2"
    assert q["options"]["C"] == "P(C)=0.3"
    assert q["options"]["D"] == "P(D)=0.4"
    assert q["correct_option"] == "A"

    # Verify RPC payload passed to database also has the full preserved options
    rpc_calls = mock_db.rpc.call_args_list
    assert len(rpc_calls) >= 1
    rpc_payload = rpc_calls[0][0][1]["p_questions"]
    assert rpc_payload[0]["options"]["A"] == "P(A)=0.1"
    assert rpc_payload[0]["options"]["B"] == "P(B)=0.2"
    assert rpc_payload[0]["options"]["C"] == "P(C)=0.3"
    assert rpc_payload[0]["options"]["D"] == "P(D)=0.4"


def test_rpc_success_readback_failure_is_not_reported_as_write_failure(
    client, mock_auth_user, auth_headers
):
    """Verify that when RPC succeeds but subsequent read-back query fails, it returns 200 with fallback payload."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    # RPC succeeds
    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 1}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    # Read-back fails with exception
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.side_effect = (
        Exception("Post-RPC query read timeout")
    )

    raw_text = """### Câu 1: Giải phương trình $x + 1 = 2$
A. $x = 1$
B. $x = 2$
C. $x = 3$
D. $x = 4$

BẢNG ĐÁP ÁN:
1.A
"""
    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    # Must NOT fail with 500 error since RPC transaction succeeded
    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 1
    assert res_data["questions"][0]["correct_key"] == "A"


def test_sync_supabase_execute_does_not_block_event_loop():
    """Verify replace_source_questions_atomically runs sync db calls in asyncio.to_thread."""
    import asyncio

    from app.models.schemas import ParsedQuestionItem
    from app.services.parser.ingestion_service import replace_source_questions_atomically

    mock_db = MagicMock()
    mock_rpc = MagicMock()
    mock_rpc.data = {"success": True, "count": 1}
    mock_db.rpc.return_value.execute.return_value = mock_rpc

    mock_q_res = MagicMock()
    mock_q_res.data = [{"question_no": 1, "raw_content": "Stem", "options": {}, "correct_key": "A"}]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    q = ParsedQuestionItem(question_no=1, raw_content="Stem", options={"A": "1", "B": "2", "C": "3", "D": "4"}, correct_key="A")

    with patch("asyncio.to_thread", wraps=asyncio.to_thread) as spy_to_thread:
        res = asyncio.run(replace_source_questions_atomically(mock_db, "session-1", [q]))
        assert res is not None
        assert spy_to_thread.called


def test_uncertain_commit_is_not_reported_as_confirmed_rollback(
    client, mock_auth_user, auth_headers
):
    """Verify that network timeout during RPC raises 504 GATEWAY_TIMEOUT with uncertain commit message."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    # RPC fails with network timeout
    mock_db.rpc.return_value.execute.side_effect = TimeoutError("Connection to Supabase timed out")

    raw_text = """### Câu 1: Tìm x
A. 1
B. 2
C. 3
D. 4
BẢNG ĐÁP ÁN: 1.A"""

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 504
    assert "uncertain" in response.json()["detail"].lower()


def test_parser_or_validator_error_prevents_database_rpc(
    client, mock_auth_user, auth_headers
):
    """Verify that any validation error prevents database RPC mutation."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    # Missing answer in text
    raw_text = """### Câu 1: Không có đáp án
A. 1
B. 2
C. 3
D. 4"""

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 422
    assert not mock_db.rpc.called


def test_ingest_file_markdown_with_bom_preserves_all_questions_and_answers(
    client, mock_auth_user, auth_headers
):
    """[P2 Regression Test] Uploading Markdown file with UTF-8 BOM preserves all questions, stem, options, and answers."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "A"},
        {"id": "q-2", "question_no": 2, "correct_key": "B"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    # 2-question Markdown file with UTF-8 BOM (\xef\xbb\xbf) and directly marked answers (*A*, *B*)
    md_bom_content = (
        b"\xef\xbb\xbf### C\xc3\xa2u 1: Gi\xc3\xa1 tr\xe1\xbb\x8b c\xe1\xbb\xa7a x trong ph\xc6\xb0\xc6\xa1ng tr\xc3\xacnh x + 1 = 2 l\xc3\xa0:\n"
        b"*A*. 1\n"
        b"B. 2\n"
        b"C. 3\n"
        b"D. 4\n\n"
        b"### C\xc3\xa2u 2: Gi\xc3\xa1 tr\xe1\xbb\x8b c\xe1\xbb\xa7a x trong ph\xc6\xb0\xc6\xa1ng tr\xc3\xacnh x - 2 = 0 l\xc3\xa0:\n"
        b"A. 1\n"
        b"*B*. 2\n"
        b"C. 3\n"
        b"D. 4\n"
    )

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "exam_with_bom.md",
                    md_bom_content,
                    "text/markdown",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 2

    # Verify Storage upload was called
    mock_db.storage.from_.return_value.upload.assert_called_once()

    # Verify RPC payload: both question 1 and question 2 received with full integrity
    assert mock_db.rpc.called
    rpc_call_args = mock_db.rpc.call_args[0][1]["p_questions"]
    assert len(rpc_call_args) == 2

    # Validate Question 1
    q1 = rpc_call_args[0]
    assert q1["question_no"] == 1
    assert "x + 1 = 2" in q1["raw_content"]
    assert q1["options"]["A"] == "1"
    assert q1["options"]["B"] == "2"
    assert q1["options"]["C"] == "3"
    assert q1["options"]["D"] == "4"
    assert q1["correct_key"] == "A"

    # Validate Question 2
    q2 = rpc_call_args[1]
    assert q2["question_no"] == 2
    assert "x - 2 = 0" in q2["raw_content"]
    assert q2["options"]["A"] == "1"
    assert q2["options"]["B"] == "2"
    assert q2["options"]["C"] == "3"
    assert q2["options"]["D"] == "4"
    assert q2["correct_key"] == "B"


def test_ingest_text_with_bom_char_preserves_all_questions(
    client, mock_auth_user, auth_headers
):
    """[P2 Regression Test] Ingesting text with BOM character preserves all questions."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "C"},
        {"id": "q-2", "question_no": 2, "correct_key": "D"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    raw_text = (
        "\ufeff### Câu 1: Tìm nguyên hàm của f(x) = 0?\n"
        "A. x  B. 1  *C*. C  D. 0\n\n"
        "### Câu 2: Tìm đạo hàm của f(x) = x?\n"
        "A. 0  B. x  C. 2  *D*. 1\n"
    )

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-text",
            json={"raw_text": raw_text},
            headers=auth_headers,
        )

    assert response.status_code == 200
    assert response.json()["total_questions"] == 2
    rpc_call_args = mock_db.rpc.call_args[0][1]["p_questions"]
    assert len(rpc_call_args) == 2
    assert rpc_call_args[0]["question_no"] == 1
    assert rpc_call_args[0]["correct_key"] == "C"
    assert rpc_call_args[1]["question_no"] == 2
    assert rpc_call_args[1]["correct_key"] == "D"


def test_ingest_file_latex_with_bom_preserves_all_questions_and_answers(
    client, mock_auth_user, auth_headers
):
    """[P2 Regression Test] Uploading LaTeX file with UTF-8 BOM preserves all questions, stem, options, and answers."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "A"},
        {"id": "q-2", "question_no": 2, "correct_key": "C"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    tex_bom_content = (
        b"\xef\xbb\xbf\\documentclass{article}\n\\begin{document}\n"
        b"\\textbf{C\xc3\xa2u 1:} Cho ph\xc6\xb0\xc6\xa1ng tr\xc3\xacnh $x^2 = 4$.\n"
        b"A. $x = \\pm 2$  B. $x = 2$  C. $x = -2$  D. $x = 0$\n\n"
        b"\\textbf{C\xc3\xa2u 2:} \xc4\x90\xe1\xba\xa1o h\xc3\xa0m c\xe1\xbb\xa7a $y = x^3$ l\xc3\xa0:\n"
        b"A. $x^2$  B. $2x$  C. $3x^2$  D. $3x$\n\n"
        b"\\textbf{B\xe1\xba\xa2NG \xc4\x90\xc3\x81P \xc3\x81N:}\n1.A 2.C\n"
        b"\\end{document}\n"
    )

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "exam_with_bom.tex",
                    tex_bom_content,
                    "application/x-tex",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 2

    # Verify RPC payload
    assert mock_db.rpc.called
    rpc_call_args = mock_db.rpc.call_args[0][1]["p_questions"]
    assert len(rpc_call_args) == 2
    assert rpc_call_args[0]["question_no"] == 1
    assert rpc_call_args[0]["correct_key"] == "A"
    assert rpc_call_args[1]["question_no"] == 2
    assert rpc_call_args[1]["correct_key"] == "C"


def test_ingest_file_markdown_heading_bold_and_closing_paren(
    client, mock_auth_user, auth_headers
):
    """Uploading Markdown with heading+bold format and closing paren question header succeeds."""
    mock_db = MagicMock()
    mock_ownership = MagicMock()
    mock_ownership.data = [{"user_id": "00000000-0000-0000-0000-000000000001"}]
    mock_status = MagicMock()
    mock_status.data = [{"status": "draft"}]

    mock_db.table.return_value.select.return_value.eq.return_value.execute.side_effect = [
        mock_ownership,
        mock_status,
    ]

    mock_db.storage.from_.return_value.upload.return_value = MagicMock()

    mock_rpc_exec = MagicMock()
    mock_rpc_exec.data = {"success": True, "count": 2}
    mock_db.rpc.return_value.execute.return_value = mock_rpc_exec

    mock_q_res = MagicMock()
    mock_q_res.data = [
        {"id": "q-1", "question_no": 1, "correct_key": "B"},
        {"id": "q-2", "question_no": 2, "correct_key": "D"},
    ]
    mock_db.table.return_value.select.return_value.eq.return_value.order.return_value.execute.return_value = mock_q_res

    md_content = (
        "### **Câu 1:** Đồ thị hàm số bậc ba?\n"
        "A. 1\n*B*. 2\nC. 3\nD. 4\n\n"
        "  Câu 2) Tìm cực trị của hàm số?\n"
        "A. 1\nB. 2\nC. 3\n*D*. 4\n"
    ).encode("utf-8")

    with (
        patch("app.api.dependencies.get_supabase_client", return_value=mock_auth_user),
        patch("app.api.dependencies.get_authenticated_client", return_value=mock_db),
    ):
        response = client.post(
            f"/api/sessions/{MOCK_SESSION_ID}/ingest-file",
            files={
                "file": (
                    "exam_heading_bold.md",
                    md_content,
                    "text/markdown",
                )
            },
            headers=auth_headers,
        )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_questions"] == 2
    rpc_call_args = mock_db.rpc.call_args[0][1]["p_questions"]
    assert len(rpc_call_args) == 2
    assert rpc_call_args[0]["question_no"] == 1
    assert rpc_call_args[0]["correct_key"] == "B"
    assert rpc_call_args[1]["question_no"] == 2
    assert rpc_call_args[1]["correct_key"] == "D"





